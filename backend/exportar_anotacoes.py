import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pymongo import DESCENDING

from app.database import classification_jobs_collection, persuasion_results_collection


ANNOTATION_HEADERS = [
    "job_id", "record_id", "approach", "experiments_tag", "chunk_index",
    "orador", "texto_completo", "status", "superclasse",
    "explicacao_superclasse", "evidencia_superclasse", "offsets_superclasse",
    "subclasse", "explicacao_subclasse", "evidencia_subclasse",
    "offsets_subclasse", "sem_classificacao_explicacao", "classified_at",
]

ERROR_HEADERS = [
    "job_id", "record_id", "approach", "experiments_tag", "chunk_index",
    "orador", "tipo_erro", "mensagem_detalhada", "data_erro",
]

SEGMENT_HEADERS = [
    "job_id", "record_id", "chunk_index", "segmento", "inicio", "fim",
    "maximo_paragrafos", "maximo_caracteres",
]

TASK_HEADERS = [
    "task_id", "status", "job_id", "audiencia_id", "fala_numero",
    "paragrafo_numero", "orador", "inicio", "fim", "contexto_anterior",
    "texto_alvo", "contexto_posterior", "categorias_modelo",
    "evidencias_modelo", "explicacoes_modelo", "possui_evidencia_nao_literal",
    "alertas", "modelo", "versao_prompt", "prompt_system", "prompt_user",
    "resposta_original_json", "estrato_multilabel", "anotadores_necessarios",
    "ordem_amostra",
]

STRATIFICATION_HEADERS = [
    "classe", "disponiveis", "selecionados", "minimo_solicitado", "cobertura_atingida",
]

VALIDATION_HEADERS = [
    "timestamp", "task_id", "avaliador", "decisao_classificacao",
    "categorias_corrigidas", "evidencia_correta", "trecho_evidencia_corrigido",
    "comentario", "duracao_segundos",
]

CONFIG_ROWS = [
    {"grupo": "decisao_classificacao", "valor": "Correta", "descricao": "Aceitar todas as categorias propostas."},
    {"grupo": "decisao_classificacao", "valor": "Incorreta", "descricao": "Corrigir uma ou mais categorias."},
    {"grupo": "decisao_classificacao", "valor": "Incerta", "descricao": "Encaminhar para adjudicação."},
    {"grupo": "evidencia_correta", "valor": "Sim", "descricao": "Os trechos sustentam as categorias."},
    {"grupo": "evidencia_correta", "valor": "Não", "descricao": "Informar o trecho corrigido."},
    {"grupo": "evidencia_correta", "valor": "Não se aplica", "descricao": "Usar quando a classificação for Nenhuma."},
    *[
        {"grupo": "categoria", "valor": value, "descricao": "Categoria permitida."}
        for value in [
            "ataque_a_reputacao", "justificativa", "simplificacao",
            "distracao", "chamada", "linguagem_manipulativa", "nenhuma",
        ]
    ],
]

INSTRUCTION_ROWS = [
    {"etapa": 1, "instrucao": "Leia o contexto anterior e posterior apenas para compreender o texto-alvo."},
    {"etapa": 2, "instrucao": "Avalie conjuntamente todas as categorias propostas para o parágrafo."},
    {"etapa": 3, "instrucao": "Marque a classificação como Correta, Incorreta ou Incerta."},
    {"etapa": 4, "instrucao": "Quando incorreta, informe as categorias corrigidas separadas por vírgula."},
    {"etapa": 5, "instrucao": "Valide as evidências. Se necessário, copie o menor trecho literal correto."},
    {"etapa": 6, "instrucao": "O Apps Script gravará cada envio na aba Validacoes usando o task_id."},
]


def _iso(value):
    return value.isoformat() if isinstance(value, datetime) else value


def _evidence_text(evidence):
    return "\n".join(item.get("text", "") for item in evidence or [])


def _evidence_offsets(evidence):
    return "\n".join(
        f'{item.get("start")}-{item.get("end")}' for item in evidence or []
    )


def _flat_spans(item):
    texts = item.get("text_spans", [])
    offsets = item.get("span_offsets", [])
    return "\n".join(texts), "\n".join(
        f'{offset.get("start")}-{offset.get("end")}' for offset in offsets
    )


def _base_row(document, job):
    return {
        "job_id": job["job_id"],
        "record_id": job["record_id"],
        "approach": job["approach"],
        "experiments_tag": job["experiments_tag"],
        "chunk_index": document.get("chunk_index"),
        "orador": document.get("speaker_name"),
        "texto_completo": document.get("text"),
        "classified_at": _iso(document.get("classified_at")),
    }


def build_export_rows(documents, job):
    annotations = []
    errors = []
    segments = []

    for document in documents:
        base = _base_row(document, job)
        error = document.get("classification_error")
        if error:
            errors.append({
                **{key: base.get(key) for key in ERROR_HEADERS},
                "tipo_erro": error.get("type"),
                "mensagem_detalhada": error.get("message"),
                "data_erro": _iso(error.get("occurred_at")),
            })
            continue

        classification = document.get("classification") or {}
        segmentation = classification.get("segmentation") or {}
        for segment in segmentation.get("segments", []):
            segments.append({
                "job_id": base["job_id"],
                "record_id": base["record_id"],
                "chunk_index": base["chunk_index"],
                "segmento": segment.get("index", 0) + 1,
                "inicio": segment.get("start"),
                "fim": segment.get("end"),
                "maximo_paragrafos": segmentation.get("max_paragraphs"),
                "maximo_caracteres": segmentation.get("max_chars"),
            })

        grouped = classification.get("classifications")
        if grouped:
            for group in grouped:
                subclasses = group.get("subclasses") or [{}]
                for subclass in subclasses:
                    annotations.append({
                        **base,
                        "status": "classificado",
                        "superclasse": group.get("superclass"),
                        "explicacao_superclasse": group.get("explanation"),
                        "evidencia_superclasse": _evidence_text(group.get("evidence")),
                        "offsets_superclasse": _evidence_offsets(group.get("evidence")),
                        "subclasse": subclass.get("subclass"),
                        "explicacao_subclasse": subclass.get("explanation"),
                        "evidencia_subclasse": _evidence_text(subclass.get("evidence")),
                        "offsets_subclasse": _evidence_offsets(subclass.get("evidence")),
                        "sem_classificacao_explicacao": None,
                    })
            continue

        flat_subclasses = classification.get("subclass_classifications")
        if flat_subclasses is None:
            flat_subclasses = classification.get("annotations")
        if flat_subclasses:
            superclasses = {
                item.get("superclass"): item
                for item in classification.get("superclass_classifications", [])
            }
            for subclass in flat_subclasses:
                superclass = superclasses.get(subclass.get("superclass"), {})
                superclass_text, superclass_offsets = _flat_spans(superclass)
                subclass_text, subclass_offsets = _flat_spans(subclass)
                annotations.append({
                    **base,
                    "status": "classificado",
                    "superclasse": subclass.get("superclass"),
                    "explicacao_superclasse": superclass.get("explanation"),
                    "evidencia_superclasse": superclass_text,
                    "offsets_superclasse": superclass_offsets,
                    "subclasse": subclass.get("subclass"),
                    "explicacao_subclasse": subclass.get("justification"),
                    "evidencia_subclasse": subclass_text,
                    "offsets_subclasse": subclass_offsets,
                    "sem_classificacao_explicacao": None,
                })
            continue

        label = classification.get("label")
        if label and label != "nenhuma":
            annotations.append({
                **base,
                "status": "classificado",
                "superclasse": label,
                "explicacao_superclasse": classification.get("explanation"),
                "evidencia_superclasse": classification.get("excerpt"),
                "offsets_superclasse": f'{classification.get("start")}-{classification.get("end")}',
                "subclasse": None,
                "explicacao_subclasse": None,
                "evidencia_subclasse": None,
                "offsets_subclasse": None,
                "sem_classificacao_explicacao": None,
            })
            continue

        annotations.append({
            **base,
            "status": "sem_classificacao",
            "superclasse": None,
            "explicacao_superclasse": None,
            "evidencia_superclasse": None,
            "offsets_superclasse": None,
            "subclasse": None,
            "explicacao_subclasse": None,
            "evidencia_subclasse": None,
            "offsets_subclasse": None,
            "sem_classificacao_explicacao": (
                classification.get("no_classification_explanation")
                or classification.get("explanation")
            ),
        })
    return annotations, errors, segments


def build_validation_tasks(documents, job):
    """Build one human-review task per classified paragraph/call."""
    tasks = []
    for document in documents:
        if document.get("classification_error"):
            continue
        classification = document.get("classification") or {}
        grouped = classification.get("classifications") or []
        groups_by_paragraph = {}
        for group in grouped:
            groups_by_paragraph.setdefault(group.get("paragraph_index", 0), []).append(group)

        audit = document.get("audit") or {}
        executions = audit.get("executions") or []
        if not executions:
            executions = [{
                "paragraph_index": 0,
                "start": 0,
                "end": len(document.get("text") or ""),
                "target_text": document.get("text") or "",
                "prompt": audit.get("prompt") or {},
                "response": audit.get("response") or classification,
                "validation_warnings": audit.get("validation_warnings") or [],
            }]

        for position, execution in enumerate(executions):
            paragraph_index = execution.get("paragraph_index", position)
            paragraph_groups = groups_by_paragraph.get(paragraph_index, [])
            categories = list(dict.fromkeys(
                group.get("superclass") for group in paragraph_groups if group.get("superclass")
            )) or ["nenhuma"]
            evidences = []
            explanations = []
            warnings = list(execution.get("validation_warnings") or [])
            unreliable = False
            for group in paragraph_groups:
                explanation = group.get("explanation")
                if explanation:
                    explanations.append(f'{group.get("superclass")}: {explanation}')
                for evidence in group.get("evidence") or []:
                    evidences.append(f'{group.get("superclass")}: {evidence.get("text", "")}')
                    if evidence.get("reliable") is False:
                        unreliable = True
                    if evidence.get("warning"):
                        warnings.append(evidence["warning"])
                if group.get("evidence_reliable") is False:
                    unreliable = True
                warnings.extend(group.get("evidence_warnings") or [])

            previous = executions[position - 1].get("target_text", "") if position > 0 else ""
            following = (
                executions[position + 1].get("target_text", "")
                if position + 1 < len(executions) else ""
            )
            prompt = execution.get("prompt") or {}
            task_id = (
                f'{job["job_id"]}:{document.get("chunk_index", 0)}:{paragraph_index}'
            )
            tasks.append({
                "task_id": task_id,
                "status": "Pendente",
                "job_id": job["job_id"],
                "audiencia_id": job["record_id"],
                "fala_numero": document.get("chunk_index", 0) + 1,
                "paragrafo_numero": paragraph_index + 1,
                "orador": document.get("speaker_name"),
                "inicio": execution.get("start"),
                "fim": execution.get("end"),
                "contexto_anterior": previous,
                "texto_alvo": execution.get("target_text", ""),
                "contexto_posterior": following,
                "categorias_modelo": ", ".join(categories),
                "evidencias_modelo": "\n".join(evidences),
                "explicacoes_modelo": "\n".join(explanations),
                "possui_evidencia_nao_literal": unreliable,
                "alertas": "\n".join(dict.fromkeys(filter(None, warnings))),
                "modelo": document.get("model") or job.get("model"),
                "versao_prompt": document.get("prompt_version") or job.get("prompt_version"),
                "prompt_system": prompt.get("system"),
                "prompt_user": prompt.get("user"),
                "resposta_original_json": json.dumps(
                    execution.get("response") or {}, ensure_ascii=False, sort_keys=True
                ),
                "estrato_multilabel": ", ".join(sorted(categories)),
                "anotadores_necessarios": 1,
                "ordem_amostra": None,
            })
    return tasks


def _write_sheet(workbook, title, headers, rows, widths):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    header_fill = PatternFill("solid", fgColor="17384D")
    sheet.row_dimensions[1].height = 28
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = widths.get(header, 18)
    for row_index, row in enumerate(rows, 2):
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(row_index, column, row.get(header))
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return sheet


def export_workbook(job, documents, output_path):
    annotations, errors, segments = build_export_rows(documents, job)
    tasks = build_validation_tasks(documents, job)
    workbook = Workbook()
    workbook.remove(workbook.active)
    wide = {
        "texto_completo": 70, "explicacao_superclasse": 45,
        "explicacao_subclasse": 45, "sem_classificacao_explicacao": 45,
        "evidencia_superclasse": 35, "evidencia_subclasse": 35,
        "mensagem_detalhada": 80,
    }
    task_widths = {
        "task_id": 48, "job_id": 38, "orador": 24,
        "contexto_anterior": 60, "texto_alvo": 70, "contexto_posterior": 60,
        "categorias_modelo": 28, "evidencias_modelo": 55,
        "explicacoes_modelo": 55, "alertas": 55, "versao_prompt": 42,
        "prompt_system": 75, "prompt_user": 75, "resposta_original_json": 75,
    }
    tasks_sheet = _write_sheet(workbook, "Tarefas", TASK_HEADERS, tasks, task_widths)
    tasks_sheet.freeze_panes = "G2"
    tasks_sheet.sheet_properties.tabColor = "17384D"
    for row_index in range(2, len(tasks) + 2):
        tasks_sheet.row_dimensions[row_index].height = 84
    for column in ["C", "H", "I", "R", "S", "T", "U", "V"]:
        tasks_sheet.column_dimensions[column].hidden = True
    tasks_sheet.conditional_formatting.add(
        f"A2:V{max(2, len(tasks) + 1)}",
        FormulaRule(
            formula=["$P2=TRUE"],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )

    validations_sheet = _write_sheet(
        workbook,
        "Validacoes",
        VALIDATION_HEADERS,
        [],
        {
            "timestamp": 22, "task_id": 48, "avaliador": 28,
            "decisao_classificacao": 24, "categorias_corrigidas": 40,
            "evidencia_correta": 22, "trecho_evidencia_corrigido": 60,
            "comentario": 60, "duracao_segundos": 18,
        },
    )
    validations_sheet.sheet_properties.tabColor = "2F855A"
    decision_validation = DataValidation(
        type="list", formula1='"Correta,Incorreta,Incerta"', allow_blank=False
    )
    evidence_validation = DataValidation(
        type="list", formula1='"Sim,Não,Não se aplica"', allow_blank=False
    )
    validations_sheet.add_data_validation(decision_validation)
    validations_sheet.add_data_validation(evidence_validation)
    decision_validation.add("D2:D2000")
    evidence_validation.add("F2:F2000")

    config_sheet = _write_sheet(
        workbook, "Config", ["grupo", "valor", "descricao"], CONFIG_ROWS,
        {"grupo": 28, "valor": 30, "descricao": 60},
    )
    config_sheet.sheet_properties.tabColor = "7C8B94"
    instructions_sheet = _write_sheet(
        workbook, "Instrucoes", ["etapa", "instrucao"], INSTRUCTION_ROWS,
        {"etapa": 10, "instrucao": 110},
    )
    instructions_sheet.sheet_properties.tabColor = "A0AEC0"
    annotations_sheet = _write_sheet(
        workbook, "Anotacoes", ANNOTATION_HEADERS, annotations, wide
    )
    for row_index in range(2, len(annotations) + 2):
        annotations_sheet.row_dimensions[row_index].height = 72
    _write_sheet(workbook, "Erros", ERROR_HEADERS, errors, wide)
    _write_sheet(workbook, "Segmentacao", SEGMENT_HEADERS, segments, wide)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(annotations), len(errors), len(segments)


def export_stratified_validation_workbook(tasks, summary_rows, output_path):
    """Export a sampled task set for the Google Apps Script annotation app."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    task_widths = {
        "task_id": 48, "job_id": 38, "orador": 24,
        "contexto_anterior": 60, "texto_alvo": 70, "contexto_posterior": 60,
        "categorias_modelo": 28, "evidencias_modelo": 55,
        "explicacoes_modelo": 55, "alertas": 55, "versao_prompt": 42,
        "prompt_system": 75, "prompt_user": 75, "resposta_original_json": 75,
        "estrato_multilabel": 38, "anotadores_necessarios": 24,
        "ordem_amostra": 18,
    }
    tasks_sheet = _write_sheet(workbook, "Tarefas", TASK_HEADERS, tasks, task_widths)
    tasks_sheet.freeze_panes = "G2"
    tasks_sheet.sheet_properties.tabColor = "17384D"
    for row_index in range(2, len(tasks) + 2):
        tasks_sheet.row_dimensions[row_index].height = 84
    for header in [
        "job_id", "inicio", "fim", "modelo", "versao_prompt", "prompt_system",
        "prompt_user", "resposta_original_json",
    ]:
        tasks_sheet.column_dimensions[
            get_column_letter(TASK_HEADERS.index(header) + 1)
        ].hidden = True
    tasks_sheet.conditional_formatting.add(
        f"A2:Y{max(2, len(tasks) + 1)}",
        FormulaRule(
            formula=["$P2=TRUE"],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )

    _write_sheet(
        workbook, "Validacoes", VALIDATION_HEADERS, [],
        {"timestamp": 22, "task_id": 48, "avaliador": 28, "comentario": 60},
    )
    _write_sheet(
        workbook, "Config", ["grupo", "valor", "descricao"], CONFIG_ROWS,
        {"grupo": 28, "valor": 30, "descricao": 60},
    )
    _write_sheet(
        workbook, "Instrucoes", ["etapa", "instrucao"], INSTRUCTION_ROWS,
        {"etapa": 10, "instrucao": 110},
    )
    summary_sheet = _write_sheet(
        workbook, "Estratificacao", STRATIFICATION_HEADERS, summary_rows,
        {"classe": 30, "disponiveis": 16, "selecionados": 16,
         "minimo_solicitado": 20, "cobertura_atingida": 20},
    )
    summary_sheet.sheet_properties.tabColor = "805AD5"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(tasks)


def _default_output(record_id, experiments_tag):
    safe_tag = re.sub(r"[^a-zA-Z0-9._-]+", "-", experiments_tag).strip("-") or "experimento"
    return Path(f"anotacoes_audiencia_{record_id}_{safe_tag}.xlsx")


def main():
    parser = argparse.ArgumentParser(description="Exporta anotações de uma classificação para XLSX.")
    parser.add_argument("--id", type=int, required=True, help="ID da audiência")
    selector = parser.add_mutually_exclusive_group()
    selector.add_argument("--experiments-tag", help="Nome/tag do experimento")
    selector.add_argument("--job-id", help="ID exato da execução")
    selector.add_argument(
        "--latest", action="store_true",
        help="Usa a execução de abordagem mais recente da audiência (padrão).",
    )
    parser.add_argument("--output", type=Path, help="Caminho do arquivo XLSX")
    args = parser.parse_args()

    query = {"record_id": args.id, "approach": {"$exists": True}}
    if args.experiments_tag:
        query["experiments_tag"] = args.experiments_tag
    if args.job_id:
        query["job_id"] = args.job_id
    job = classification_jobs_collection.find_one(
        query,
        sort=[("created_at", DESCENDING)],
    )
    if not job:
        parser.error(f"Nenhuma classificação encontrada para audiência {args.id}.")
    documents = list(
        persuasion_results_collection.find({"job_id": job["job_id"]}, {"_id": 0})
        .sort("chunk_index", 1)
    )
    output = args.output or _default_output(args.id, job["experiments_tag"])
    annotation_count, error_count, segment_count = export_workbook(job, documents, output)
    task_count = len(build_validation_tasks(documents, job))
    print(f"Execução selecionada: {job['job_id']} ({job['approach']})")
    print(f"Arquivo criado: {output.resolve()}")
    print(f"Anotações: {annotation_count}; erros: {error_count}; segmentos: {segment_count}")
    print(f"Tarefas de validação humana: {task_count}")


if __name__ == "__main__":
    main()
