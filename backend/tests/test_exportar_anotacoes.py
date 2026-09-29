import tempfile
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from exportar_amostra_estratificada import (
    optimize_balanced_multilabel_tasks,
    stratify_by_cardinality_and_class,
    stratify_multilabel_tasks,
)
from exportar_anotacoes import build_export_rows, build_validation_tasks, export_workbook


class ExportAnnotationsTests(unittest.TestCase):
    def setUp(self):
        self.job = {
            "job_id": "job-123",
            "record_id": 42,
            "approach": "hierarchical_multiclass_segmented_v3",
            "experiments_tag": "experimento-a",
        }

    def test_builds_annotation_error_and_segmentation_rows(self):
        documents = [{
            "chunk_index": 3,
            "speaker_name": "Maria",
            "text": "Precisamos agir agora.",
            "classified_at": "2026-09-05T12:00:00Z",
            "classification": {
                "classifications": [{
                    "superclass": "chamada",
                    "explanation": "Convoca uma ação.",
                    "evidence": [{"text": "agir agora", "start": 11, "end": 21}],
                    "subclasses": [{
                        "subclass": "apelo_ao_tempo",
                        "explanation": "Indica urgência.",
                        "evidence": [{"text": "agir agora", "start": 11, "end": 21}],
                    }],
                }],
                "segmentation": {
                    "max_paragraphs": 3,
                    "max_chars": 4000,
                    "segments": [{"index": 0, "start": 0, "end": 22}],
                },
            },
        }, {
            "chunk_index": 4,
            "speaker_name": "João",
            "text": "Texto inválido.",
            "classification_error": {
                "type": "RuntimeError",
                "message": "Falha no segmento 1",
                "occurred_at": "2026-09-05T12:01:00Z",
            },
        }]

        annotations, errors, segments = build_export_rows(documents, self.job)

        self.assertEqual(annotations[0]["subclasse"], "apelo_ao_tempo")
        self.assertEqual(annotations[0]["offsets_subclasse"], "11-21")
        self.assertEqual(errors[0]["mensagem_detalhada"], "Falha no segmento 1")
        self.assertEqual(segments[0]["segmento"], 1)

    def test_builds_one_validation_task_per_audited_paragraph(self):
        documents = [{
            "chunk_index": 0,
            "speaker_name": "Ana",
            "text": "Primeiro.\nSegundo.",
            "classification": {"classifications": [{
                "superclass": "chamada",
                "paragraph_index": 1,
                "explanation": "Convoca ação.",
                "evidence": [{
                    "text": "Segundo", "start": 10, "end": 17,
                    "reliable": False, "warning": "Não literal.",
                }],
            }]},
            "audit": {"executions": [{
                "paragraph_index": 0, "start": 0, "end": 9,
                "target_text": "Primeiro.", "prompt": {},
                "response": {"anotacoes": [{"categoria": "Nenhuma"}]},
            }, {
                "paragraph_index": 1, "start": 10, "end": 18,
                "target_text": "Segundo.", "prompt": {"system": "s", "user": "u"},
                "response": {"anotacoes": [{"categoria": "Call"}]},
            }]},
        }]

        tasks = build_validation_tasks(documents, self.job)

        self.assertEqual(len(tasks), 2)
        self.assertEqual(tasks[0]["categorias_modelo"], "nenhuma")
        self.assertEqual(tasks[1]["categorias_modelo"], "chamada")
        self.assertTrue(tasks[1]["possui_evidencia_nao_literal"])
        self.assertEqual(tasks[1]["contexto_anterior"], "Primeiro.")

    def test_exports_a_readable_validation_workbook(self):
        documents = [{
            "chunk_index": 0,
            "speaker_name": "Ana",
            "text": "Bom dia.",
            "classification": {
                "classifications": [],
                "no_classification_explanation": "Somente uma saudação.",
            },
        }]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "anotacoes.xlsx"
            counts = export_workbook(self.job, documents, output)
            workbook = load_workbook(output, read_only=True)

            self.assertEqual(counts, (1, 0, 0))
            self.assertEqual(workbook.sheetnames, [
                "Tarefas", "Validacoes", "Config", "Instrucoes",
                "Anotacoes", "Erros", "Segmentacao",
            ])
            self.assertEqual(workbook["Anotacoes"]["H2"].value, "sem_classificacao")
            self.assertEqual(workbook["Tarefas"]["B2"].value, "Pendente")
            self.assertEqual(workbook["Validacoes"]["B1"].value, "task_id")

    def test_validation_tasks_skip_classification_errors(self):
        documents = [{
            "chunk_index": 0,
            "text": "Falhou.",
            "classification_error": {"type": "APIConnectionError", "message": "Connection error."},
        }]
        self.assertEqual(build_validation_tasks(documents, self.job), [])

    def test_multilabel_stratification_counts_one_task_for_each_of_its_labels(self):
        labels = [
            "ataque_a_reputacao", "justificativa", "simplificacao", "distracao",
            "chamada", "linguagem_manipulativa", "nenhuma",
        ]
        tasks = []
        for label in labels:
            for index in range(2):
                tasks.append({
                    "task_id": f"{label}-{index}",
                    "categorias_modelo": label,
                })
        tasks.append({
            "task_id": "multi",
            "categorias_modelo": "ataque_a_reputacao, chamada",
        })

        selected, summary = stratify_multilabel_tasks(tasks, minimum_per_class=2, seed=7)

        coverage = {row["classe"]: row["selecionados"] for row in summary}
        self.assertTrue(all(value >= 2 for value in coverage.values()))
        self.assertTrue(all(task["anotadores_necessarios"] == 2 for task in selected))
        self.assertEqual([task["ordem_amostra"] for task in selected], list(range(1, len(selected) + 1)))

    def test_balanced_optimizer_respects_cardinality_targets_and_redundancy(self):
        labels = [
            "ataque_a_reputacao", "justificativa", "simplificacao", "distracao",
            "chamada", "linguagem_manipulativa",
        ]
        tasks = []
        for cardinality in range(1, 7):
            for index in range(12):
                chosen = labels[index % len(labels):] + labels[:index % len(labels)]
                values = chosen[:cardinality]
                tasks.append({
                    "task_id": f"{cardinality}-{index}",
                    "categorias_modelo": ", ".join(values),
                })
        for index in range(12):
            tasks.append({"task_id": f"none-{index}", "categorias_modelo": "nenhuma"})

        selected, _ = optimize_balanced_multilabel_tasks(
            tasks,
            cardinality_targets={1: 8, 2: 4, 3: 3, 4: 2, 5: 2, 6: 1},
            minimum_per_class=2,
            seed=7,
            iterations=2_000,
        )

        cardinalities = Counter(len(task["categorias_modelo"].split(",")) for task in selected)
        self.assertEqual(cardinalities, Counter({1: 8, 2: 4, 3: 3, 4: 2, 5: 2, 6: 1}))
        self.assertTrue(all(task["anotadores_necessarios"] == 2 for task in selected))

    def test_cardinality_first_stratification_uses_empty_stratum_and_distinct_quotas(self):
        labels = [
            "ataque_a_reputacao", "justificativa", "simplificacao", "distracao",
            "chamada", "linguagem_manipulativa",
        ]
        tasks = []
        for cardinality in range(1, 7):
            for index in range(8):
                rotated = labels[index % 6:] + labels[:index % 6]
                tasks.append({
                    "task_id": f"{cardinality}-{index}",
                    "categorias_modelo": ", ".join(rotated[:cardinality]),
                })
        tasks.extend(
            {"task_id": f"none-{index}", "categorias_modelo": "nenhuma"}
            for index in range(8)
        )
        selected, _ = stratify_by_cardinality_and_class(
            tasks, quota_per_cardinality=4, empty_quota=4, seed=7, iterations=500
        )
        strata = Counter(task["estrato_multilabel"] for task in selected)
        self.assertEqual(strata["vazia"], 4)
        self.assertTrue(all(strata[f"{cardinality}_labels"] == 4 for cardinality in range(1, 7)))
        self.assertTrue(all(task["anotadores_necessarios"] == 2 for task in selected))


if __name__ == "__main__":
    unittest.main()
